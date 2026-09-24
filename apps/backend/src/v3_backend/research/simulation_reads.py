"""Consistent snapshot reads of paper-account tables, curves and exports."""
from pathlib import Path
import csv
import json
import math
from .storage import identifier
from .simulation_accounts import get_record,revision


def iter_rows(db,account,table,check):
    if table=='ownership':
        for row in account.get('ownership',[]):check();yield row
        return
    kind={'cash_flows':'simulation_cash_flow','binding_events':'simulation_binding_event'}.get(table)
    if kind:
        for record in db.execute('SELECT body FROM records WHERE kind=? ORDER BY rowid',(kind,)):
            check();row=json.loads(record[0])
            if row.get('accountId')==account['id']:yield row
        if table=='cash_flows':return
    for record in db.execute('SELECT body FROM records WHERE kind=? AND id LIKE ? ORDER BY id',('simulation_day',account['id']+':%')):
        check();day=json.loads(record[0])
        for row in day['tables'].get(table,[]):yield row


def dispatch(store,method,params,check=lambda:None):
    from .simulation import TABLES
    portable=store.project_store(params.get('projectId'))
    with portable.connect() as db:
        db.execute('BEGIN');check();account=get_record(db,'simulation',params['accountId'])
        if account.get('projectId')!=params.get('projectId'):raise ValueError('账户不属于此范围')
        if method=='simulation.table':
            table=params['table']
            if table not in TABLES:raise ValueError('未知模拟表')
            offset=max(0,int(params.get('offset',0)));limit=min(500,max(1,int(params.get('limit',200))));rows=[];total=0;columns=[]
            for row in iter_rows(db,account,table,check):
                if params.get('symbol') and row.get('symbol')!=params['symbol']:continue
                columns=list(dict.fromkeys(columns+list(row)))
                if offset<=total<offset+limit:rows.append(row)
                total+=1
            return dict(name=table,columns=columns,rows=rows,total=total,offset=offset,limit=limit)
        if method=='simulation.accounts.curve':
            maximum=max(2,min(2000,int(params.get('maxPoints',600))))
            total=sum(1 for _ in iter_rows(db,account,'portfolio',check))
            indices={round(i*(total-1)/(maximum-1)) for i in range(min(total,maximum))} if total>maximum else set(range(total))
            rows=[]
            for index,row in enumerate(iter_rows(db,account,'portfolio',check)):
                if index in indices:rows.append(dict(date=row['date'],unitNav=row.get('unitNav'),nav=row.get('account')))
            return dict(rows=rows,total=total)
        if method!='simulation.accounts.export':raise ValueError('未知账户读取操作')
        revision(account,params.get('expectedRevision'));format=params.get('format')
        if format not in {'csv','xlsx'}:raise ValueError('账户导出只支持CSV/XLSX')
        table=params.get('table')
        if table is not None and table not in TABLES:raise ValueError('未知模拟表')
        if format=='csv' and table is None:raise ValueError('CSV导出请选择数据表')
        selected=[table] if table else list(TABLES)
        folder=Path(store.project(params.get('projectId'))['path'])/'exports';folder.mkdir(parents=True,exist_ok=True)
        path=folder/('account-'+account['id']+'-'+identifier()+'.'+format);temporary=path.with_suffix('.writing.'+format)
        book=None
        def cell(value):
            if isinstance(value,(dict,list)):return json.dumps(value,ensure_ascii=False,allow_nan=False)
            return value
        try:
            if format=='xlsx':
                from openpyxl import Workbook
                from openpyxl.cell import WriteOnlyCell
                book=Workbook(write_only=True)
            for name in selected:
                columns=[]
                for row in iter_rows(db,account,name,check):columns=list(dict.fromkeys(columns+list(row)))
                if format=='csv':
                    with temporary.open('w',encoding='utf-8-sig',newline='') as stream:
                        writer=csv.writer(stream);writer.writerow(columns)
                        for row in iter_rows(db,account,name,check):writer.writerow([cell(row.get(key)) for key in columns])
                else:
                    sheet=book.create_sheet(name);sheet.append(columns);count=0;part=0
                    for row in iter_rows(db,account,name,check):
                        if count==1048575:part+=1;sheet=book.create_sheet(name+'_'+str(part));sheet.append(columns);count=0
                        values=[]
                        for key in columns:
                            value=cell(row.get(key))
                            if isinstance(value,str):
                                if len(value)>32767:raise ValueError('Excel单元格过长，请改用CSV保留完整内容')
                                value=WriteOnlyCell(sheet,value=value);value.data_type='s'
                            values.append(value)
                        sheet.append(values);count+=1
            check()
            if book:book.save(temporary)
            check();temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)
            if book:
                for sheet in book.worksheets:
                    writer=getattr(sheet,'_writer',None)
                    if writer is not None:
                        if not sheet.closed:sheet.close()
                        if Path(writer.out).exists():writer.cleanup()
                book.close()
        return dict(path=str(path))
