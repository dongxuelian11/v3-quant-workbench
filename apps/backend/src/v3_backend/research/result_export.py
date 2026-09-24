"""Export full experiment artifacts without materializing all tables together."""
import re
import uuid
from pathlib import Path
import pandas as pd
import pyarrow.parquet as pq
from .table_reader import batches


def export(store, project_id, experiment, output, format, table=None, *, check_cancel=None):
    if check_cancel: check_cancel()
    artifacts = [a for a in experiment['artifacts'] if a['type'] == 'parquet']
    if format not in {'csv', 'xlsx'}:
        raise ValueError('只支持 CSV/XLSX 导出')
    if table and not any(a['name'] == table for a in artifacts):
        raise ValueError('实验中没有此数据表')
    path = Path(output) / (experiment['id'] + '-export-' + uuid.uuid4().hex + '.' + format)
    temporary = path.with_suffix('.' + uuid.uuid4().hex + '.writing.' + format)
    book = None

    def frames(artifact):
        if check_cancel: check_cancel()
        if artifact is None:
            yield pd.DataFrame(list(experiment['metrics'].items()), columns=['metric', 'value'])
        else:
            with pq.ParquetFile(store.artifact_path(project_id, artifact)) as source:
                if source.metadata.num_rows == 0:
                    yield pd.DataFrame(columns=source.schema_arrow.names)
                else:
                    for batch in batches(source, check_cancel=check_cancel, batch_size=8192):
                        frame = batch.to_pandas()
                        if check_cancel: check_cancel()
                        yield frame

    try:
        if format == 'csv':
            selected = next((a for a in artifacts if a['name'] == table), artifacts[0] if artifacts else None)
            with temporary.open('w', encoding='utf-8-sig', newline='') as stream:
                for i, frame in enumerate(frames(selected)):
                    if check_cancel: check_cancel()
                    frame.to_csv(stream, index=False, header=i == 0)
        else:
            from openpyxl import Workbook
            book = Workbook(write_only=True)
            selected_artifacts = [a for a in artifacts if a['name'] == table] if table else artifacts
            for index, artifact in enumerate(selected_artifacts or [None]):
                name = re.sub(r'[\\/*?:\[\]]', '_', artifact['name'] if artifact else 'metrics')
                sheet, count, chunk = None, 0, 0
                for frame in frames(artifact):
                    if sheet is None:
                        sheet = book.create_sheet(f'{index}_{chunk}_{name}'[:31])
                        sheet.append(list(frame.columns))
                    for row in frame.itertuples(index=False, name=None):
                        if count % 256 == 0 and check_cancel: check_cancel()
                        if count == 1048575:
                            chunk += 1
                            sheet = book.create_sheet(f'{index}_{chunk}_{name}'[:31])
                            sheet.append(list(frame.columns))
                            count = 0
                        if any(isinstance(value,str) and len(value)>32767 for value in row):
                            raise ValueError('Excel 单元格最多保存32767字符；此成果包含更长的完整诊断，请使用 CSV 导出，避免内容截断')
                        sheet.append([None if pd.api.types.is_scalar(value) and pd.isna(value) else value for value in row])
                        count += 1
            if check_cancel: check_cancel()
            book.save(temporary)
        if check_cancel: check_cancel()
        # Commit point: a later cancel cannot retract a completed export.
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
        if book is not None:
            # write-only worksheets own additional XML tempfiles until save succeeds.
            for sheet in book.worksheets:
                writer = getattr(sheet, '_writer', None)
                if writer is not None:
                    if not sheet.closed:
                        sheet.close()
                    if Path(writer.out).exists():
                        writer.cleanup()
            book.close()
    return {'path': str(path)}


def export_comparison(store, project_id, experiment_ids, output, format, *, check_cancel=None):
    import json
    from .storage import read_json
    if format not in {'csv', 'xlsx'}:
        raise ValueError('比较表支持 CSV/XLSX 导出')
    if not isinstance(experiment_ids, list) or len(set(experiment_ids)) < 2:
        raise ValueError('请选择至少两个不同实验')
    rows = []
    for key in dict.fromkeys(experiment_ids):
        if check_cancel: check_cancel()
        experiment = store.experiment(project_id, key)
        root = Path(store.project(project_id)['path']) / '.research/runs' / key
        snapshot = read_json(root / 'project.json', {})
        params = experiment.get('parameters', {})
        row = dict(experimentId=key, name=experiment['name'], kind=experiment['kind'],
            createdAt=experiment.get('createdAt'), projectId=project_id,
            startDate=params.get('startDate') or snapshot.get('startDate'),
            endDate=params.get('endDate') or snapshot.get('endDate'),
            universe=json.dumps(snapshot.get('universe'), ensure_ascii=False),
            parameters=json.dumps(params, ensure_ascii=False))
        row.update({'metric.' + str(key): value if value is None or isinstance(value, (str, int, float, bool))
                    else json.dumps(value, ensure_ascii=False) for key, value in experiment.get('metrics', {}).items()})
        rows.append(row)
    path = Path(output) / ('comparison-' + uuid.uuid4().hex + '.' + format)
    temporary = path.with_suffix('.writing.' + format)
    try:
        frame = pd.DataFrame(rows)
        if format == 'csv': frame.to_csv(temporary, index=False, encoding='utf-8-sig')
        else: frame.to_excel(temporary, index=False, sheet_name='实验比较', engine='openpyxl')
        if check_cancel: check_cancel()
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    return {'path': str(path)}
