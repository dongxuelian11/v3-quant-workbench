import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from v3_backend.research.result_export import export


class ResultExport(unittest.TestCase):
    def test_excel_selected_table_excludes_other_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            artifacts = []
            for name, values in [('trades',[1]),('rule_diagnostics',[10,20,30])]:
                path = root/(name+'.parquet')
                pd.DataFrame({'value':values}).to_parquet(path,index=False)
                artifacts.append(dict(name=name,type='parquet',path=str(path)))
            experiment = dict(id='result',metrics={},artifacts=artifacts)
            store = SimpleNamespace(artifact_path=lambda project, value:Path(value['path']))
            selected = export(store,None,experiment,root,'xlsx','rule_diagnostics')
            book = load_workbook(selected['path'],read_only=True)
            try:
                self.assertEqual(book.sheetnames,['0_0_rule_diagnostics'])
                self.assertEqual(list(book.active.values),[('value',),(10,),(20,),(30,)])
            finally:book.close()
            complete = export(store,None,experiment,root,'xlsx')
            book = load_workbook(complete['path'],read_only=True)
            try:
                self.assertEqual(book.sheetnames,['0_0_trades','1_0_rule_diagnostics'])
            finally:book.close()

    def test_csv_and_excel_keep_all_rows_across_batches(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            frame = pd.DataFrame({'date': pd.date_range('2000-01-01', periods=8200),
                                  'value': range(8200)})
            artifact = root / 'input.parquet'
            frame.to_parquet(artifact, index=False)
            experiment = dict(id='result', metrics={}, artifacts=[dict(name='daily', type='parquet', path=str(artifact))])
            store = SimpleNamespace(artifact_path=lambda project, value: Path(value['path']))
            with patch.object(pd, 'read_parquet', side_effect=AssertionError('export must stream')):
                csv = export(store, None, experiment, root, 'csv')
                xlsx = export(store, None, experiment, root, 'xlsx')
            actual = pd.read_csv(csv['path'])
            self.assertEqual(len(actual), 8200)
            self.assertEqual(actual.value.iloc[-1], 8199)
            book = load_workbook(xlsx['path'], read_only=True)
            try:
                rows = list(book.active.values)
                self.assertEqual(len(rows), 8201)
                self.assertEqual(rows[-1][-1], 8199)
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
