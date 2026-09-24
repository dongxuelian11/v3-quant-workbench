"""Small frozen JSON fixtures; no computation, market-table reads or downloads."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from v3_backend.research.server import Service
from v3_backend.research.storage import write_json
from v3_backend.research import comparison


class ComparisonTest(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory()
        self.root=Path(self.temporary.name)
        self.service=Service(self.root/'app')
        self.store=self.service.store
        self.project=self.store.create_project(self.root/'one','one')
        self.other=self.store.create_project(self.root/'two','two')
        self.details=dict(engine='pyqlib',accountingBasis='raw_shares_cash',signalTiming='previous close -> next open',
            strategyPriceBasis='raw',annualizationDays=252,annualizedReturnMethod='arithmetic')
        self.identity=dict(sourceRoot='recorded-source/data',sourceFiles=[['prices.parquet',100,1024]],
            localRoot='recorded-project',localMembershipIndustryFiles=[],
            slice=dict(startDate='2025-01-01',endDate='2025-03-31',symbols=['SH600000']))

    def tearDown(self):
        self.service.close()
        self.temporary.cleanup()

    def saved(self,name,created='2025-04-01T00:00:00Z',project=None,params=None,identity=True,snapshot=True,
              strategy='default',details=None,source=None,universe=None):
        project=project or self.project
        frozen=dict(path=project['path'],startDate='2025-01-01',endDate='2025-03-31',
            universe=universe or dict(source='manual',symbols=['SH600000'],excludeST=False,minListingDays=0))
        values=dict(startDate='2025-01-01',endDate='2025-03-31',benchmark='csi300',topN=3)
        values.update(params or {})
        data=deepcopy(self.details if details is None else details)
        data['dataContext']=[dict(name='prices',rows=60,startDate='2025-01-01',endDate='2025-03-31')]
        folder=Path(project['path'])/'.research/runs'/name
        reference=dict(version=1,status='available',path=f'.research/runs/{name}/inputs',createdAt=created,
                       files=[{'path':'data/prices.parquet'}])
        if source:reference['sourceExperimentId']=source
        if identity:frozen['inputCacheIdentity']=deepcopy(self.identity if identity is True else identity)
        if snapshot:
            write_json(folder/'inputs/snapshot.json',dict(project=frozen,parameters=values,reference=reference))
            payload=folder/'inputs/data/prices.parquet'
            payload.parent.mkdir()
            payload.write_text('fixture payload; comparison must not read it')
        write_json(folder/'project.json',frozen)
        write_json(folder/'details.json',data)
        experiment=dict(id=name,projectId=project['id'],name=name,kind='backtest.run',createdAt=created,
            parameters=values,metrics={'total_return':.1},summary='fixture',artifacts=[])
        if strategy is not None:experiment['strategyId']=strategy
        if snapshot:experiment['inputSnapshot']=reference
        self.store.save_experiment(project['id'],experiment)
        return dict(projectId=project['id'],experimentId=name)

    def compared(self,refs,**extra):
        with patch.object(pd,'read_parquet',side_effect=AssertionError('metadata only')):
            return self.service.request('experiments.compare',dict(experiments=refs,**extra))

    def test_controlled_configuration_range_and_frozen_context(self):
        a=self.saved('A')
        b=self.saved('B',params={'topN':5})
        results=self.compared([a,b])
        comp=results[1]['details']['comparison']
        self.assertEqual(comp['comparability']['status'],'controlled_change')
        self.assertEqual(comp['configurationChanges'][0]['field'],'topN')
        self.assertEqual(comp['inputEvidence']['status'],'same_recorded_version')
        self.assertIn('不证明',comp['inputEvidence']['message'])
        c=self.saved('C',params={'startDate':'2025-02-01','benchmark':'csi500'},
            universe=dict(source='manual',symbols=['SH600001'],excludeST=False,minListingDays=0))
        comp=self.compared([a,c])[1]['details']['comparison']
        self.assertEqual(comp['comparability']['status'],'incomparable')
        self.assertTrue({'startDate','benchmark','universe.symbols'} <= {r['field'] for r in comp['rangeChanges']})
        # Only the project path may come from today's project record, never its semantics.
        current=self.store.project(self.project['id'])
        current.update(startDate='2030-01-01',universe={'source':'manual','symbols':['SZ000001']})
        self.store.save_project(current)
        self.assertEqual(comp,self.compared([a,c])[1]['details']['comparison'])
        swapped=self.compared([a,b],baselineRef=b)
        self.assertEqual(swapped[0]['details']['comparison']['baselineRef'],b)
        with self.assertRaisesRegex(ValueError,'基准'):
            self.compared([a,b],baselineRef=c)

    def test_same_coverage_is_not_same_content_and_missing_evidence_wins(self):
        a=self.saved('A',identity=False)
        b=self.saved('B',identity=False,params={'topN':5})
        # Unexamined market payloads can differ with exactly the same coverage.
        for reference,value in [(a,10),(b,99)]:
            path=Path(self.project['path'])/'.research/runs'/reference['experimentId']/'inputs/data'
            path.mkdir(exist_ok=True)
            (path/'prices.parquet').write_text(str(value))
        comp=self.compared([a,b])[1]['details']['comparison']
        self.assertEqual(comp['inputEvidence']['status'],'unknown')
        self.assertEqual(comp['comparability']['status'],'unknown')
        old=self.saved('old',snapshot=False,identity=False,details={},params={'startDate':'2024-01-01'})
        comp=self.compared([a,old])[1]['details']['comparison']
        self.assertEqual(comp['comparability']['status'],'unknown')
        self.assertTrue(any(r['code']=='context_changed' for r in comp['comparability']['reasons']))
        self.assertEqual(comp['inputEvidence']['current']['snapshotStatus'],'missing')

    def test_reproduction_and_recorded_source_versions(self):
        a=self.saved('A',identity=False)
        b=self.saved('B',identity=False,source='A')
        comp=self.compared([a,b])[1]['details']['comparison']
        self.assertEqual(comp['inputEvidence']['basis'],'snapshot_reference')
        self.assertEqual(comp['inputEvidence']['status'],'same_recorded_version')
        c=self.saved('C')
        identity=deepcopy(self.identity);identity['sourceFiles'][0][1]=200
        d=self.saved('D',identity=identity)
        comp=self.compared([c,d])[1]['details']['comparison']
        self.assertEqual(comp['inputEvidence']['status'],'different_recorded_version')
        self.assertEqual(comp['comparability']['status'],'unknown')
        identity=deepcopy(self.identity);identity['sourceRoot']='another/directory'
        e=self.saved('E',identity=identity)
        self.assertEqual(self.compared([c,e])[1]['details']['comparison']['inputEvidence']['status'],'unknown')

    def test_cross_project_same_ids_and_consistent_exports(self):
        a=self.saved('same')
        b=self.saved('same',project=self.other,params={'benchmark':'csi500'})
        details=self.compared([a,b])
        self.assertEqual(len(details),2)
        expected={r['details']['comparison']['ref']['projectId']:r['details']['comparison'] for r in details}
        csv=self.service.request('exports.create',dict(experiments=[a,b],format='csv'))
        xlsx=self.service.request('exports.create',dict(experiments=[a,b],format='xlsx'))
        rows=pd.read_csv(csv['path'],keep_default_na=False)
        self.assertEqual(set(rows.projectId),{self.project['id'],self.other['id']})
        for _,row in rows[rows.field.eq('comparison')].iterrows():
            self.assertEqual(json.loads(row.value),expected[row.projectId])
        book=load_workbook(xlsx['path'],read_only=True)
        try:
            self.assertEqual(book.sheetnames,['实验身份','配置变化','范围与口径','输入证据','可比性','指标'])
            values=list(book['实验身份'].values);columns=values.pop(0)
            for row in map(lambda v:dict(zip(columns,v)),values):
                if row['field']=='comparison':self.assertEqual(json.loads(row['value']),expected[row['projectId']])
        finally:book.close()
        # Legacy callers keep projectId + experimentIds; readable identity/date are retained.
        c=self.saved('C',params={'topN':10})
        exported=self.service.request('exports.create',dict(projectId=self.project['id'],experimentIds=['same','C'],format='csv'))
        self.assertTrue(Path(exported['path']).exists())
        ordinary=self.service.request('exports.create',dict(projectId=self.project['id'],experimentId='C',format='csv'))
        self.assertEqual(list(pd.read_csv(ordinary['path']).columns),['metric','value'])

    def test_previous_read_cancellation_and_missing_snapshot_metadata(self):
        a=self.saved('A')
        b=self.saved('B','2025-04-02T00:00:00Z')
        operation=self.service.read_operations.begin('experiments.previousComparison',{'readId':'previous'})
        self.service.read_operations.cancel('previous')
        with self.assertRaises(InterruptedError):
            self.service.request('experiments.previousComparison',b,read_operation=operation)
        # Registered references with missing metadata do not become evidence of matching contents.
        metadata=Path(self.project['path'])/'.research/runs/B/inputs/snapshot.json'
        metadata.rename(metadata.with_suffix('.missing'))
        comp=self.compared([a,b])[1]['details']['comparison']
        self.assertEqual(comp['inputEvidence']['status'],'unknown')
        self.assertEqual(comp['inputEvidence']['current']['snapshotStatus'],'unavailable')
        invalid=self.saved('invalid-time',created='not-recorded')
        with self.assertRaisesRegex(ValueError,'创建时间'):
            self.service.request('experiments.previousComparison',invalid)

    def test_supplemental_files_without_versions_and_inline_code_changes(self):
        for field,extension in [('codePath','py'),('dataPath','parquet')]:
            with self.subTest(field=field):
                a=self.saved('A-'+field,params={'customFactors':[{'id':'external',field:'copied/a.'+extension}]})
                b=self.saved('B-'+field,params={'customFactors':[{'id':'external',field:'copied/b.'+extension}]})
                for reference,value in [(a,'old code or values'),(b,'new code or values')]:
                    root=Path(self.project['path'])/'.research/runs'/reference['experimentId']/'inputs'
                    (root/('extra.'+extension)).write_text(value)
                comp=self.compared([a,b])[1]['details']['comparison']
                self.assertEqual(comp['comparability']['status'],'unknown')
                self.assertEqual(comp['inputEvidence']['status'],'unknown')
                self.assertIn(field,comp['inputEvidence']['message'])
                self.assertIn('未比对附加输入',comp['inputEvidence']['message'])
                self.assertEqual(comp['configurationChanges'],[]) # Copy locations alone are not strategy changes.
        a=self.saved('inlineA',params={'code':'scores = scores * 2','expression':'$close'})
        b=self.saved('inlineB',params={'code':'scores = scores * 3','expression':'$open'})
        comp=self.compared([a,b])[1]['details']['comparison']
        self.assertEqual(comp['comparability']['status'],'controlled_change')
        self.assertEqual({r['field'] for r in comp['configurationChanges']},{'code','expression'})

    def test_missing_listed_payload_is_unavailable_without_content_reads(self):
        a=self.saved('A')
        b=self.saved('B')
        payload=Path(self.project['path'])/'.research/runs/B/inputs/data/prices.parquet'
        payload.rename(payload.with_suffix('.missing'))
        comp=self.compared([a,b])[1]['details']['comparison']
        self.assertEqual(comp['inputEvidence']['current']['snapshotStatus'],'unavailable')
        self.assertEqual(comp['inputEvidence']['status'],'unknown')
        self.assertEqual(comp['comparability']['status'],'unknown')
        self.assertIn('data/prices.parquet',comp['inputEvidence']['message'])

    def test_previous_is_strict_deterministic_successful_and_lightweight(self):
        none=self.saved('first','2025-04-01T00:00:00Z',strategy=None)
        self.assertIsNone(self.service.request('experiments.previousComparison',none)['previousRef'])
        self.saved('wrong-strategy','2025-04-02T00:00:00Z',strategy='other')
        a=self.saved('A','2025-04-02T00:00:00Z')
        b=self.saved('B','2025-04-02T00:00:00Z')
        partial=self.saved('partial','2025-04-03T00:00:00Z',details={**self.details,'status':'partial','incomplete':True})
        cancelled=self.saved('cancelled','2025-04-04T00:00:00Z')
        self.store.project_store(self.project['id']).put('job',dict(id='cancelled',status='cancelled',experimentId='cancelled'),self.project['id'])
        failed=self.saved('failed','2025-04-04T12:00:00Z')
        write_json(Path(self.project['path'])/'.research/runs/failed/result.json',{'runtimeStatus':'failed'})
        wrong=self.saved('wrong-kind','2025-04-04T13:00:00Z')
        experiment=self.store.experiment(self.project['id'],'wrong-kind');experiment['kind']='model.train'
        self.store.save_experiment(self.project['id'],experiment)
        current=self.saved('current','2025-04-05T00:00:00Z',strategy=None)
        self.saved('equal-time','2025-04-05T08:00:00+08:00')
        self.saved('future','2025-04-06T00:00:00Z')
        with patch.object(self.service,'experiment_details',side_effect=AssertionError('no large previews')):
            result=self.service.request('experiments.previousComparison',current)
        self.assertEqual(result['previousRef'],b)  # Equal prior timestamps use experiment id deterministically.
        self.assertEqual(result['comparison']['baselineRef'],b)
        self.assertEqual(result['comparison'],self.compared([b,current])[1]['details']['comparison'])


if __name__ == '__main__':unittest.main()
