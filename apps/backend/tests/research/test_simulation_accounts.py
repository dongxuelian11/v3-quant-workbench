import tempfile
import unittest
from pathlib import Path
from copy import deepcopy
from v3_backend.research.server import Service
from v3_backend.research import simulation_accounts as accounts


class SharedAccountTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.service=Service(Path(self.tmp.name)/'app');self.store=self.service.store
        self.account=self.call('create',name='账户',capital=1000,startDate='2025-01-02')
    def tearDown(self):
        self.service.close();self.tmp.cleanup()
    def call(self,op,**params):return self.service.request('simulation.accounts.'+op,params)
    def test_cash_flow_idempotency_conflicts_and_persistence(self):
        self.assertEqual(self.account['bindings'],[])
        p=dict(accountId=self.account['id'],expectedRevision=0,requestId='deposit',direction='deposit',amount=500)
        result=self.call('cashFlows.create',**p);a=result['account']
        self.assertEqual((a['cash'],a['nav'],a['units'],a['unitNav']),(1500,1500,1500,1))
        self.assertEqual(self.call('cashFlows.create',**p),result)
        with self.assertRaisesRegex(ValueError,'不同内容'):self.call('cashFlows.create',**{**p,'amount':100})
        with self.assertRaisesRegex(ValueError,'版本'):self.call('save',accountId=a['id'],expectedRevision=0,name='stale')
        with self.assertRaisesRegex(ValueError,'现金不足'):self.call('cashFlows.create',**{**p,'requestId':'too-much','expectedRevision':1,'direction':'withdraw','amount':1501})
        final=self.call('cashFlows.create',**{**p,'requestId':'withdraw','expectedRevision':1,'direction':'withdraw','amount':200})['account']
        self.assertEqual((final['cash'],final['unitNav'],final['netContributions']),(1300,1,1300))
        table=self.service.request('simulation.table',dict(accountId=a['id'],table='cash_flows',offset=1,limit=1))
        self.assertEqual(table['total'],2);self.assertEqual(len(table['rows']),1)
        self.assertEqual(self.store.project_store(None).get('simulation',a['id']),final)
    def test_binding_versions_manual_transfer_and_branch_checkpoint(self):
        project=self.store.create_project(Path(self.tmp.name)/'project','源')
        project['settings'].setdefault('backtest',{})['dailyCode']='old code';self.store.save_project(project)
        a=self.call('bindings.save',accountId=self.account['id'],expectedRevision=0,source=dict(kind='strategy',projectId=project['id'],strategyId='default'),allocation=.5)
        binding=a['bindings'][0];old=binding['currentVersionId']
        strategy=self.store.project_store(project['id']).get('strategy','default');strategy['settings']['backtest']['dailyCode']='new code';self.store.project_store(project['id']).put('strategy',strategy,project['id'])
        preview=self.call('bindings.preview',accountId=a['id'],bindingId=binding['id'])
        a=self.call('bindings.adopt',accountId=a['id'],expectedRevision=1,bindingId=binding['id'],expectedSourceVersion=preview['sourceVersion'])
        self.assertEqual(len(a['bindings'][0]['versions']),2)
        self.assertEqual(a['bindings'][0]['versions'][0]['snapshot']['project']['settings']['backtest']['dailyCode'],'old code')
        self.assertNotEqual(a['bindings'][0]['currentVersionId'],old)
        # A completed checkpoint represents a full account state, not today's state.
        a['asOfDate']='2025-01-02';a['state']['asOfDate']='2025-01-02'
        a['state']['holdings']={'SH600000':dict(quantity=100,sellableQuantity=100)}
        a['ownership']=[dict(id='owner',symbol='SH600000',bindingId=binding['id'],entryVersionId=old,management='rules',quantity=100,sellableQuantity=100,pendingQuantity=0,costBasis=None)]
        with self.store.project_store(None).connect() as db:
            accounts.put_record(db,'simulation',a);accounts.checkpoint(db,a)
        args=dict(accountId=a['id'],expectedRevision=2,fromDate='2025-01-02',name='历史分支',requestId='branch')
        branch=self.call('branch',**args)
        self.assertNotEqual(branch['id'],a['id']);self.assertEqual(branch['ownership'][0]['entryVersionId'],old)
        self.assertEqual(self.call('branch',**args),branch)
        manual=self.call('ownership.transferToManual',accountId=branch['id'],expectedRevision=0,ownershipIds=['owner'])
        self.assertEqual(manual['ownership'][0]['management'],'manual');self.assertEqual(manual['state']['holdings'],a['state']['holdings'])
        self.assertEqual(self.call('get',accountId=a['id'])['ownership'][0]['management'],'rules')
        with self.assertRaisesRegex(ValueError,'不存在'):self.call('branch',**{**args,'fromDate':'2024-12-31','requestId':'missing'})
    def test_legacy_import_preserves_source_and_unknown_history(self):
        project=self.store.create_project(Path(self.tmp.name)/'legacy','旧账户')
        legacy=self.call('create',projectId=project['id'],strategyId='default',capital=1000,startDate='2025-01-02')
        legacy['state']['holdings']={'SH600000':dict(quantity=100,sellableQuantity=100)}
        self.store.project_store(project['id']).put('simulation',legacy,project['id'])
        p=dict(source=dict(projectId=project['id'],accountId=legacy['id']),expectedRevision=0,requestId='import')
        imported=self.call('importLegacy',**p)
        self.assertIsNone(imported['projectId']);self.assertIsNone(imported['unitNav']);self.assertEqual(imported['valuationStatus'],'unavailable')
        self.assertEqual(imported['ownership'][0]['management'],'manual');self.assertEqual(imported['bindings'],[])
        self.assertEqual(self.call('importLegacy',**p),imported)
        self.assertEqual(self.call('get',projectId=project['id'],accountId=legacy['id']),legacy)
        with self.assertRaisesRegex(ValueError,'版本'):self.call('importLegacy',**{**p,'requestId':'stale','expectedRevision':99})
        with self.assertRaisesRegex(ValueError,'估值'):self.call('cashFlows.create',accountId=imported['id'],expectedRevision=0,requestId='unknown',direction='deposit',amount=10)
        self.assertEqual(len(self.call('list')),3)

    def test_preview_race_partial_manual_transfer_and_strict_ownership(self):
        from unittest.mock import patch
        project=self.store.create_project(Path(self.tmp.name)/'version-source','源')
        a=self.call('bindings.save',accountId=self.account['id'],expectedRevision=0,source=dict(kind='strategy',projectId=project['id'],strategyId='default'),allocation=.5)
        binding=a['bindings'][0];preview=self.call('bindings.preview',accountId=a['id'],bindingId=binding['id'])
        changed={**preview,'sourceVersion':'changed'}
        with patch.object(accounts,'preview_source',side_effect=[preview,changed]):
            with self.assertRaisesRegex(ValueError,'重新预览'):self.call('bindings.adopt',accountId=a['id'],expectedRevision=a['revision'],bindingId=binding['id'],expectedSourceVersion=preview['sourceVersion'])
        a['state']['holdings']={'SH600000':dict(quantity=200,sellableQuantity=200)}
        a['ownership']=[dict(id=key,symbol='SH600000',bindingId=binding['id'],entryVersionId=binding['currentVersionId'],management='rules',quantity=100,sellableQuantity=100,pendingQuantity=0,costBasis=None) for key in ('a','b')]
        a['pendingDecision']=dict(date='2025-01-02',intents=[dict(ownershipId=key,targetQuantity=0,costs={'x':1},versionId='fixed') for key in ('a','b')])
        self.store.project_store(None).put('simulation',a)
        result=self.call('ownership.transferToManual',accountId=a['id'],expectedRevision=a['revision'],ownershipIds=['a'])
        self.assertEqual(result['pendingDecision']['intents'],[a['pendingDecision']['intents'][1]])
        invalid=deepcopy(result);invalid['ownership'][0]['quantity']=-1
        with self.assertRaises(ValueError):accounts.validate_ownership(invalid)
        invalid=deepcopy(result);invalid['state']['holdings']['SH600000']['quantity']=1000000000;invalid['ownership'][0]['quantity']=999999899
        with self.assertRaises(ValueError):accounts.validate_ownership(invalid)

    def test_realtime_export_curve_and_zero_balance_cashflow_phase(self):
        import csv
        from openpyxl import load_workbook
        a=self.account
        flow=self.call('cashFlows.create',accountId=a['id'],expectedRevision=0,requestId='all-out',direction='withdraw',amount=1000)
        self.assertEqual(flow['cashFlow']['effectivePhase'],'before_start')
        a=flow['account'];a=self.call('cashFlows.create',accountId=a['id'],expectedRevision=1,requestId='in-again',direction='deposit',amount=250)['account']
        self.assertEqual((a['unitNav'],a['units']),(1,250))
        exported=self.call('export',accountId=a['id'],expectedRevision=a['revision'],format='csv',table='cash_flows')
        with open(exported['path'],encoding='utf-8-sig',newline='') as stream:rows=list(csv.DictReader(stream))
        self.assertEqual(len(rows),2);self.assertTrue(all(r['effectivePhase']=='before_start' for r in rows))
        book=load_workbook(self.call('export',accountId=a['id'],expectedRevision=a['revision'],format='xlsx')['path'],read_only=True)
        try:self.assertTrue(set(('cash_flows','ownership','trade_allocations','binding_events')).issubset(book.sheetnames))
        finally:book.close()
        with self.assertRaisesRegex(ValueError,'版本'):self.call('export',accountId=a['id'],expectedRevision=0,format='xlsx')
        with self.store.project_store(None).connect() as db:
            for i in range(12):
                date='2025-01-'+str(i+1).zfill(2)
                accounts.put_record(db,'simulation_day',dict(id=a['id']+':'+date,accountId=a['id'],date=date,tables={'portfolio':[dict(date=date,account=250,unitNav=None if i==0 else 1.)]}))
        curve=self.call('curve',accountId=a['id'],maxPoints=4)
        self.assertEqual(curve['total'],12);self.assertEqual(len(curve['rows']),4)
        self.assertEqual(curve['rows'][0]['date'],'2025-01-01');self.assertIsNone(curve['rows'][0]['unitNav']);self.assertEqual(curve['rows'][-1]['date'],'2025-01-12')
